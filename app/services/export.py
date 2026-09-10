"""Serialize runs into downloadable CSV/XLSX/JSON exports (docs/TASKS_EXPORT.md).

Three query helpers (`runs_for_run`, `runs_for_prompt`, `runs_for_client`)
each select the `Run` rows for one export scope, with every relationship an
export needs already eager-loaded — the `build_*` writers below never
trigger a lazy load, even across hundreds of runs. The three writers turn
that same `list[Run]` into bytes for one format; `content` controls how
much of each run goes in (see docs/TASKS_EXPORT.md design decisions 3-4 for
why raw_payload is handled differently per format: CSV can add loose files
to its ZIP safely, XLSX cannot).
"""

import csv
import io
import json
import zipfile
from datetime import datetime, timezone
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, selectinload

from app.models import AIModel, Prompt, PromptSet, RawResponse, Run

ExportContent = Literal["answer", "raw", "full"]

EXCEL_CELL_CHAR_LIMIT = 32767
_TRUNCATE_SUFFIX = "... [truncated — use format=json for full payload]"

# A leading =, +, -, @, tab, or CR makes Excel/LibreOffice read a CSV cell as
# a formula on open (CSV/formula injection, OWASP). rendered_text/prompt_text
# can contain attacker-influenced text (a provider answer echoing scraped web
# content), so every string written to CSV is neutralized with a leading
# quote — same mitigation whether the value came from a client name or an AI
# provider response, since CSV has no way to mark a field as "definitely not
# a formula" the way a typed XLSX cell does.
_CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

RUN_COLUMNS: tuple[str, ...] = (
    "id",
    "prompt_id",
    "prompt_version",
    "is_current_version",
    "client_name",
    "prompt_set_name",
    "prompt_text",
    "market_code",
    "provider_code",
    "model_name",
    "status",
    "trigger_type",
    "started_at",
    "finished_at",
    "latency_ms",
    "error_message",
    "rendered_text",
    "has_citations",
    "token_usage",
)

CITATION_COLUMNS: tuple[str, ...] = (
    "run_id",
    "source_url",
    "source_title",
    "source_domain",
    "citation_position",
    "cited_answer_span",
)

_RUN_EAGER_LOAD = (
    selectinload(Run.raw_response).selectinload(RawResponse.citations),
    selectinload(Run.prompt).selectinload(Prompt.prompt_set).selectinload(PromptSet.client),
    selectinload(Run.model).selectinload(AIModel.provider),
    selectinload(Run.market),
)


def runs_for_run(db: Session, run_id: int) -> list[Run]:
    """The single run with this id, or an empty list if it doesn't exist.

    Only fetches data — turning an empty result into a 404 is the router's
    job, not this function's.
    """
    run = db.scalar(select(Run).where(Run.id == run_id).options(*_RUN_EAGER_LOAD))
    return [run] if run is not None else []


def runs_for_prompt(db: Session, prompt_id: int, *, all_versions: bool = False) -> list[Run]:
    """Runs for one prompt.

    `all_versions=False` (default) returns only runs against the exact
    prompt version named by `prompt_id` — the same scope `prompt_detail`
    already shows on screen (app/routers/prompts.py). `all_versions=True`
    walks the version lineage (`root_prompt_id`/`is_current_version`, same
    logic as `_version_history()` in app/routers/prompts.py) and returns
    runs against every version of that prompt. A `prompt_id` that doesn't
    exist yields an empty list either way — the router 404s separately.
    """
    prompt_ids: list[int] = [prompt_id]
    if all_versions:
        prompt = db.get(Prompt, prompt_id)
        if prompt is not None:
            root_id = prompt.root_prompt_id or prompt.id
            prompt_ids = list(
                db.scalars(select(Prompt.id).where(or_(Prompt.id == root_id, Prompt.root_prompt_id == root_id))).all()
            )
    return list(
        db.scalars(
            select(Run).where(Run.prompt_id.in_(prompt_ids)).options(*_RUN_EAGER_LOAD).order_by(Run.started_at.desc())
        ).all()
    )


def runs_for_client(db: Session, client_id: int) -> list[Run]:
    """Every run across every prompt set/prompt/version belonging to one client."""
    return list(
        db.scalars(
            select(Run)
            .join(Prompt, Run.prompt_id == Prompt.id)
            .join(PromptSet, Prompt.prompt_set_id == PromptSet.id)
            .where(PromptSet.client_id == client_id)
            .options(*_RUN_EAGER_LOAD)
            .order_by(Run.started_at.desc())
        ).all()
    )


def build_filename(scope: str, identifier: str, content: ExportContent, ext: str) -> str:
    """The download filename for one export.

    `signalmap_{scope}_{identifier}[_{content}]_{YYYYMMDD}.{ext}` — the
    content segment is only appended when it isn't the default "answer"
    tier, and the date is the export's UTC generation date (see
    docs/TASKS_EXPORT.md design decision 6).
    """
    parts = ["signalmap", scope, identifier]
    if content != "answer":
        parts.append(content)
    parts.append(datetime.now(timezone.utc).strftime("%Y%m%d"))
    return f"{'_'.join(parts)}.{ext}"


def _run_base_fields(run: Run) -> dict[str, Any]:
    """The run/prompt/model metadata every export tier includes, in every format.

    Shared by `_run_row` (CSV/XLSX) and `build_json` so the run→field mapping
    lives in exactly one place — previously each duplicated the same ~16
    attribute lookups independently, risking one being updated and the other
    forgotten on a schema change.
    """
    return {
        "id": run.id,
        "prompt_id": run.prompt_id,
        "prompt_version": run.prompt.version,
        "is_current_version": run.prompt.is_current_version,
        "client_name": run.prompt.prompt_set.client.name,
        "prompt_set_name": run.prompt.prompt_set.name,
        "prompt_text": run.prompt.text,
        "market_code": run.market.code,
        "provider_code": run.model.provider.code,
        "model_name": run.model.model_name,
        "status": run.status,
        "trigger_type": run.trigger_type,
        "started_at": run.started_at.isoformat() if run.started_at else None,
        "finished_at": run.finished_at.isoformat() if run.finished_at else None,
        "latency_ms": run.latency_ms,
        "error_message": run.error_message,
    }


def _run_row(run: Run) -> dict[str, Any]:
    """Flatten one run (+ its raw response, if any) into the `Runs`/`runs.csv` row shape.

    Always includes rendered_text/has_citations/token_usage regardless of
    `content` — CSV/XLSX only ever render the "answer" tier as tables
    (docs/TASKS_EXPORT.md design decision 3); `content` instead controls
    whether the *extra* raw-payload file/sheet gets added alongside them.
    """
    raw = run.raw_response
    row = _run_base_fields(run)
    row["rendered_text"] = raw.rendered_text if raw else None
    row["has_citations"] = raw.has_citations if raw else False
    row["token_usage"] = json.dumps(raw.token_usage) if raw and raw.token_usage else None
    return row


def _csv_safe(value: Any) -> Any:
    """Neutralize CSV/formula injection on one cell value — see `_CSV_FORMULA_PREFIXES`."""
    if isinstance(value, str) and value.startswith(_CSV_FORMULA_PREFIXES):
        return "'" + value
    return value


def _sanitize_csv_row(row: dict[str, Any]) -> dict[str, Any]:
    """Apply `_csv_safe` to every value in a row dict before it reaches `csv.DictWriter`."""
    return {key: _csv_safe(value) for key, value in row.items()}


def _xlsx_safe(value: Any) -> Any:
    """Strip characters Excel's XML format can't hold (ASCII control chars) from a string cell value.

    openpyxl raises `IllegalCharacterError` on write otherwise — provider
    output, scraped citation text, and exception messages are all realistic
    sources of a stray control character, so every string cell goes through
    this before `Worksheet.append()`.
    """
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _citation_rows(run: Run) -> list[dict[str, Any]]:
    """One row per citation on this run's raw response — empty when there is none or it has no citations."""
    raw = run.raw_response
    if raw is None:
        return []
    return [
        {
            "run_id": run.id,
            "source_url": citation.source_url,
            "source_title": citation.source_title,
            "source_domain": citation.source_domain,
            "citation_position": citation.citation_position,
            "cited_answer_span": citation.cited_answer_span,
        }
        for citation in raw.citations
    ]


def _raw_payload_text(run: Run) -> str | None:
    """Pretty-printed `raw_payload` for one run, or None when it has no raw response."""
    if run.raw_response is None:
        return None
    return json.dumps(run.raw_response.raw_payload, indent=2, default=str)


def build_csv_zip(runs: list[Run], content: ExportContent) -> bytes:
    """A ZIP with `runs.csv` + `citations.csv`, joined on `run_id`.

    `citations.csv` is always present, header-only if no run has any
    citation — downstream tooling can rely on the file existing. When
    `content` is "raw" or "full", one `raw_<run_id>.json` per run (for runs
    that have a raw response) is added alongside them — safe here because
    the CSV export is a plain ZIP with no format constraint against loose
    files, unlike XLSX (see `build_xlsx`).
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        runs_csv = io.StringIO()
        writer = csv.DictWriter(runs_csv, fieldnames=RUN_COLUMNS)
        writer.writeheader()
        for run in runs:
            writer.writerow(_sanitize_csv_row(_run_row(run)))
        zf.writestr("runs.csv", runs_csv.getvalue())

        citations_csv = io.StringIO()
        writer = csv.DictWriter(citations_csv, fieldnames=CITATION_COLUMNS)
        writer.writeheader()
        for run in runs:
            for row in _citation_rows(run):
                writer.writerow(_sanitize_csv_row(row))
        zf.writestr("citations.csv", citations_csv.getvalue())

        if content in ("raw", "full"):
            for run in runs:
                raw_text = _raw_payload_text(run)
                if raw_text is not None:
                    zf.writestr(f"raw_{run.id}.json", raw_text)

    return buffer.getvalue()


def build_xlsx(runs: list[Run], content: ExportContent) -> bytes:
    """One workbook: `Runs` + `Citations` sheets, plus a `RawPayload` sheet for raw/full content.

    `raw_payload` never becomes a loose file stuffed into the `.xlsx` zip
    container — a `.xlsx` *is* a zip, but any file outside its OOXML
    manifest makes Excel report the package as corrupted on open. The
    `RawPayload` sheet (run_id + JSON text) is the safe equivalent, with
    each cell truncated at Excel's ~32,767-character limit rather than
    raising (docs/TASKS_EXPORT.md design decision 4).
    """
    wb = Workbook()
    runs_sheet = wb.active
    runs_sheet.title = "Runs"
    runs_sheet.append(RUN_COLUMNS)
    for run in runs:
        row = _run_row(run)
        runs_sheet.append([_xlsx_safe(row[column]) for column in RUN_COLUMNS])

    citations_sheet = wb.create_sheet("Citations")
    citations_sheet.append(CITATION_COLUMNS)
    for run in runs:
        for row in _citation_rows(run):
            citations_sheet.append([_xlsx_safe(row[column]) for column in CITATION_COLUMNS])

    if content in ("raw", "full"):
        raw_sheet = wb.create_sheet("RawPayload")
        raw_sheet.append(("run_id", "raw_payload_json"))
        for run in runs:
            raw_text = _raw_payload_text(run)
            if raw_text is None:
                continue
            raw_text = _xlsx_safe(raw_text)
            if len(raw_text) > EXCEL_CELL_CHAR_LIMIT:
                raw_text = raw_text[: EXCEL_CELL_CHAR_LIMIT - len(_TRUNCATE_SUFFIX)] + _TRUNCATE_SUFFIX
            raw_sheet.append((run.id, raw_text))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_json(runs: list[Run], content: ExportContent) -> bytes:
    """A JSON array of run objects — the full-fidelity export.

    Unlike the CSV/XLSX tables (always answer-shaped, with raw data purely
    additive), JSON keys are gated by tier: `rendered_text`/`citations`
    only appear for `content in ("answer", "full")`, `raw_payload` only for
    `content in ("raw", "full")` — a "raw" export stays a minimal evidence
    dump instead of merging in unrelated answer fields.
    """
    include_answer = content in ("answer", "full")
    include_raw = content in ("raw", "full")
    entries: list[dict[str, Any]] = []
    for run in runs:
        raw = run.raw_response
        entry = _run_base_fields(run)
        if include_answer:
            entry["rendered_text"] = raw.rendered_text if raw else None
            entry["has_citations"] = raw.has_citations if raw else False
            entry["token_usage"] = raw.token_usage if raw else None
            entry["citations"] = [
                {
                    "source_url": citation.source_url,
                    "source_title": citation.source_title,
                    "source_domain": citation.source_domain,
                    "citation_position": citation.citation_position,
                    "cited_answer_span": citation.cited_answer_span,
                }
                for citation in (raw.citations if raw else [])
            ]
        if include_raw:
            entry["raw_payload"] = raw.raw_payload if raw else None
        entries.append(entry)
    return json.dumps(entries, indent=2, default=str).encode("utf-8")
